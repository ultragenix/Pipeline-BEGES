"""Ingestion config-driven des fichiers sources vers DuckDB.

Chaque source est lue selon sa configuration, convertie vers un schéma
canonique, puis écrite dans une table brute. Les lignes non exploitables sont
isolées dans ``raw.quarantine`` et chaque exécution est journalisée dans
``raw.ingestion_log``.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import hashlib
import json
import os
import sys
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from beges.config import (
    DATA_RAW_DIR,
    Company,
    Config,
    ConfigError,
    Source,
    SourceType,
    load_config,
)

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DB_PATH = Path(
    os.environ.get("BEGES_DB_PATH", ROOT / "data" / "warehouse" / "beges.duckdb")
)

TABLE_BY_TYPE = {
    "fec": "fec",
    "electricity": "electricity",
    "electricity_xlsx": "electricity",
    "fuel": "fuel",
    "refrigerants": "refrigerants",
    "flights": "flights",
    "vehicle_registry": "vehicles",
}

RAW_COLUMNS: dict[str, tuple[tuple[str, str], ...]] = {
    "fec": (
        ("entry_id", "VARCHAR"),
        ("entry_date", "DATE"),
        ("account", "VARCHAR"),
        ("account_label", "VARCHAR"),
        ("supplier_id", "VARCHAR"),
        ("supplier_label", "VARCHAR"),
        ("piece_ref", "VARCHAR"),
        ("piece_date", "DATE"),
        ("label", "VARCHAR"),
        ("debit", "DECIMAL(18,2)"),
        ("credit", "DECIMAL(18,2)"),
        ("valid_date", "DATE"),
    ),
    "electricity": (
        ("site_code", "VARCHAR"),
        ("site_code_raw", "VARCHAR"),
        ("pdl", "VARCHAR"),
        ("month", "DATE"),
        ("period_start", "DATE"),
        ("period_end", "DATE"),
        ("consumption_kwh", "DECIMAL(18,3)"),
        ("amount_ht_eur", "DECIMAL(18,2)"),
        ("reading_type", "VARCHAR"),
        ("unit_raw", "VARCHAR"),
    ),
    "fuel": (
        ("transaction_date", "DATE"),
        ("card_id", "VARCHAR"),
        ("vehicle_id", "VARCHAR"),
        ("product_raw", "VARCHAR"),
        ("quantity_l", "DECIMAL(18,3)"),
        ("amount_ttc_eur", "DECIMAL(18,2)"),
        ("station", "VARCHAR"),
    ),
    "refrigerants": (
        ("intervention_date", "DATE"),
        ("site_code", "VARCHAR"),
        ("equipment", "VARCHAR"),
        ("fluid_code", "VARCHAR"),
        ("nominal_charge_kg", "DECIMAL(18,3)"),
        ("recharge_kg", "DECIMAL(18,3)"),
    ),
    "flights": (
        ("ticket_id", "VARCHAR"),
        ("emission_date", "DATE"),
        ("departure_date", "DATE"),
        ("traveler_id", "VARCHAR"),
        ("origin_iata", "VARCHAR"),
        ("destination_iata", "VARCHAR"),
        ("trip_type", "VARCHAR"),
        ("cabin_class", "VARCHAR"),
        ("status", "VARCHAR"),
        ("amount_ttc_eur", "DECIMAL(18,2)"),
    ),
    "vehicles": (
        ("vehicle_id", "VARCHAR"),
        ("vehicle_type", "VARCHAR"),
        ("tank_capacity_l", "DECIMAL(18,3)"),
    ),
}

META_COLUMNS: tuple[tuple[str, str], ...] = (
    ("company_code", "VARCHAR NOT NULL"),
    ("source_id", "VARCHAR NOT NULL"),
    ("source_file", "VARCHAR NOT NULL"),
    ("source_sheet", "VARCHAR"),
    ("source_row", "INTEGER"),
    ("ingestion_id", "VARCHAR NOT NULL"),
    ("ingested_at", "TIMESTAMPTZ NOT NULL"),
    ("raw_row", "JSON"),
)

FIELD_KINDS: dict[str, str] = {}
for _columns in RAW_COLUMNS.values():
    for _name, _sql_type in _columns:
        if _sql_type == "DATE":
            _kind = "date"
        elif _sql_type.startswith("DECIMAL"):
            _kind = "decimal"
        else:
            _kind = "text"
        FIELD_KINDS.setdefault(_name, _kind)


class IngestError(RuntimeError):
    """Erreur bloquante d'ingestion."""


class SourceReadError(Exception):
    """Erreur de lecture d'une source, isolée sans écriture partielle."""

    def __init__(self, reason_code: str, detail: str) -> None:
        super().__init__(f"{reason_code}: {detail}")
        self.reason_code = reason_code
        self.detail = detail


@dataclass(frozen=True)
class ParsedRow:
    values: dict[str, Any]
    source_row: int
    source_sheet: str | None
    raw: dict[str, Any]


@dataclass(frozen=True)
class QuarantineRow:
    source_row: int | None
    source_sheet: str | None
    reason_code: str
    reason_detail: str | None
    raw: dict[str, Any] | None


@dataclass
class ReadResult:
    rows: list[ParsedRow] = field(default_factory=list)
    quarantine: list[QuarantineRow] = field(default_factory=list)
    rows_read: int = 0


@dataclass
class SourceLog:
    ingestion_id: str
    company_code: str
    source_id: str
    source_type: str
    source_file: str
    source_sheet: str | None
    source_sha256: str | None
    file_size_bytes: int | None
    rows_read: int
    rows_loaded: int
    rows_quarantined: int
    status: str
    error: str | None
    duration_ms: int
    started_at: datetime
    finished_at: datetime


@dataclass
class RunSummary:
    companies: int
    files: int
    rows_loaded: int
    rows_quarantined: int
    failures: int
    db_path: Path
    entries: list[SourceLog] = field(default_factory=list)


def _now() -> datetime:
    return datetime.now(UTC)


def _hash_file(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _clean_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").strip()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value: Any, separator: str) -> Decimal:
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int | float):
        number = Decimal(str(value))
    else:
        text = (
            str(value)
            .strip()
            .replace("\u00a0", "")
            .replace("\u202f", "")
            .replace(" ", "")
        )
        if separator == ",":
            if "." in text:
                raise ValueError("séparateurs décimaux ambigus")
            text = text.replace(",", ".")
        elif separator == ".":
            if "," in text:
                raise ValueError("séparateurs décimaux ambigus")
        else:
            raise ValueError(f"séparateur décimal non supporté: {separator!r}")
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError("nombre invalide") from exc
    if not number.is_finite():
        raise ValueError("nombre non fini")
    return number


def _convert_value(
    value: Any,
    kind: str,
    required: bool,
    source: Source,
) -> Any:
    text = _text(value)
    if text is None:
        if required and kind in {"date", "decimal"}:
            raise ValueError("valeur requise vide")
        return None
    if kind == "date":
        if not source.date_format:
            raise ValueError("date_format manquant")
        return datetime.strptime(text, source.date_format).date()
    if kind == "decimal":
        return _decimal(text, source.decimal)
    return text


def _resolve_site(company: Company, value: str | None) -> str | None:
    if value is None:
        return None
    if value in company.sites:
        return value
    for code, site in company.sites.items():
        if site.name == value:
            return code
    return None


def _header_positions(header: list[str], context: str) -> dict[str, int]:
    duplicates = sorted({name for name in header if name and header.count(name) > 1})
    if duplicates:
        raise SourceReadError(
            "colonne_dupliquee", f"{context}colonnes dupliquées: {', '.join(duplicates)}"
        )
    return {name: index for index, name in enumerate(header) if name}


def _validate_header(
    source: Source, positions: dict[str, int], context: str
) -> None:
    missing = sorted(
        {column for column in source.column_map.values() if column not in positions}
    )
    if missing:
        raise SourceReadError(
            "colonne_manquante",
            f"{context}colonnes absentes: {', '.join(missing)}",
        )


def _convert_row(
    company: Company,
    source: Source,
    source_type: SourceType,
    row: list[Any],
    positions: dict[str, int],
    raw: dict[str, Any],
    source_row: int,
    source_sheet: str | None,
    sheet_site: str | None,
    sheet_unit: str | None,
) -> tuple[ParsedRow | None, QuarantineRow | None]:
    mapped = {
        canonical: row[positions[column]]
        for canonical, column in source.column_map.items()
    }

    values: dict[str, Any] = {}
    for canonical in source.column_map:
        kind = FIELD_KINDS.get(canonical)
        if kind is None:
            continue
        required = canonical in source_type.required_fields
        try:
            values[canonical] = _convert_value(mapped[canonical], kind, required, source)
        except ValueError as exc:
            reason = "date_invalide" if kind == "date" else "nombre_invalide"
            return None, QuarantineRow(
                source_row,
                source_sheet,
                reason,
                f"{canonical}={mapped[canonical]!r}: {exc}",
                raw,
            )

    table = TABLE_BY_TYPE[source.type]
    table_columns = {name for name, _ in RAW_COLUMNS[table]}
    for name in table_columns:
        values.setdefault(name, None)

    if "site_code" in table_columns:
        site_raw = sheet_site if source.type == "electricity_xlsx" else values.get("site_code")
        if site_raw is None:
            return None, QuarantineRow(
                source_row,
                source_sheet,
                "valeur_requise_vide",
                "site_code",
                raw,
            )
        site_code = _resolve_site(company, str(site_raw))
        if site_code is None:
            return None, QuarantineRow(
                source_row,
                source_sheet,
                "site_inconnu",
                f"site_code={site_raw!r}",
                raw,
            )
        values["site_code"] = site_code
        if "site_code_raw" in table_columns:
            values["site_code_raw"] = str(site_raw)

    if "month" in table_columns:
        month = values.get("month")
        if month is None:
            period_start = values.get("period_start")
            if period_start is None:
                return None, QuarantineRow(
                    source_row,
                    source_sheet,
                    "mois_indetermine",
                    "month absent et period_start vide",
                    raw,
                )
            month = period_start
        values["month"] = month.replace(day=1)

    if "unit_raw" in table_columns:
        if source.type != "electricity_xlsx":
            values["unit_raw"] = None
        else:
            unit_raw = sheet_unit
            cell_unit = _text(mapped.get("unit"))
            if cell_unit:
                if sheet_unit and cell_unit.casefold() != sheet_unit.casefold():
                    return None, QuarantineRow(
                        source_row,
                        source_sheet,
                        "unite_incoherente",
                        f"ligne={cell_unit!r}, feuille={sheet_unit!r}",
                        raw,
                    )
                unit_raw = cell_unit
            if unit_raw is None:
                return None, QuarantineRow(
                    source_row,
                    source_sheet,
                    "unite_inconnue",
                    "unité absente",
                    raw,
                )
            values["unit_raw"] = unit_raw
            if unit_raw.casefold() == "mwh":
                consumption = values.get("consumption_kwh")
                if consumption is not None:
                    values["consumption_kwh"] = consumption * 1000
            elif unit_raw.casefold() != "kwh":
                return None, QuarantineRow(
                    source_row,
                    source_sheet,
                    "unite_inconnue",
                    f"unité={unit_raw!r}",
                    raw,
                )

    record = {name: values.get(name) for name in table_columns}
    return ParsedRow(record, source_row, source_sheet, raw), None


def _read_delimited(
    path: Path,
    company: Company,
    source: Source,
    source_type: SourceType,
) -> ReadResult:
    result = ReadResult()
    try:
        handle = path.open(encoding=source.encoding, newline="")
    except OSError as exc:
        raise SourceReadError("fichier_illisible", str(exc)) from exc

    with handle:
        reader = csv.reader(handle, delimiter=source.delimiter)
        try:
            raw_header = next(reader)
        except StopIteration as exc:
            raise SourceReadError("fichier_vide", "aucune ligne d'en-tête") from exc

        header = [_clean_header(value) for value in raw_header]
        if not header:
            raise SourceReadError("fichier_vide", "en-tête vide")
        positions = _header_positions(header, "")
        _validate_header(source, positions, "")

        try:
            for row in reader:
                if not any(str(value).strip() for value in row):
                    continue
                result.rows_read += 1
                raw = {
                    name: value
                    for name, value in zip(header, row, strict=False)
                    if name
                }
                if len(row) > len(header):
                    raw["_extra"] = row[len(header) :]
                if len(row) != len(header):
                    result.quarantine.append(
                        QuarantineRow(
                            reader.line_num,
                            None,
                            "ligne_malformee",
                            f"{len(header)} colonnes attendues, {len(row)} trouvées",
                            raw,
                        )
                    )
                    continue
                parsed, quarantined = _convert_row(
                    company,
                    source,
                    source_type,
                    row,
                    positions,
                    raw,
                    reader.line_num,
                    None,
                    None,
                    None,
                )
                if parsed is not None:
                    result.rows.append(parsed)
                if quarantined is not None:
                    result.quarantine.append(quarantined)
        except UnicodeDecodeError as exc:
            raise SourceReadError("encodage_invalide", str(exc)) from exc
        except csv.Error as exc:
            raise SourceReadError("ligne_malformee", str(exc)) from exc
    return result


def _read_xlsx(
    path: Path,
    company: Company,
    source: Source,
    source_type: SourceType,
) -> ReadResult:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException, ValueError) as exc:
        raise SourceReadError("fichier_illisible", str(exc)) from exc

    result = ReadResult()
    header_row = source.header_row or 1
    try:
        for sheet_name, sheet_config in source.sheets.items():
            if sheet_name not in workbook.sheetnames:
                raise SourceReadError("feuille_manquante", sheet_name)
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(min_row=header_row, values_only=True)
            try:
                raw_header = next(rows)
            except StopIteration as exc:
                raise SourceReadError(
                    "fichier_vide", f"feuille {sheet_name}: aucune ligne"
                ) from exc

            header = [_clean_header(value) for value in raw_header]
            while header and header[-1] == "":
                header.pop()
            positions = _header_positions(header, f"feuille {sheet_name}: ")
            _validate_header(source, positions, f"feuille {sheet_name}: ")

            sheet_site = _text(sheet_config.get("site_code"))
            if _resolve_site(company, sheet_site) is None:
                raise SourceReadError(
                    "site_inconnu", f"feuille {sheet_name}: site {sheet_site!r}"
                )
            sheet_unit = _text(sheet_config.get("unit"))
            if sheet_unit is None:
                raise SourceReadError(
                    "unite_inconnue", f"feuille {sheet_name}: unité absente"
                )

            for excel_row, values in enumerate(rows, start=header_row + 1):
                row = list(values[: len(header)])
                row.extend([None] * (len(header) - len(row)))
                if not any(
                    value is not None and str(value).strip() for value in row
                ):
                    continue
                result.rows_read += 1
                raw = {
                    name: value
                    for name, value in zip(header, row, strict=False)
                    if name
                }
                extra = values[len(header) :]
                if any(value is not None and str(value).strip() for value in extra):
                    raw["_extra"] = list(extra)
                    result.quarantine.append(
                        QuarantineRow(
                            excel_row,
                            sheet_name,
                            "ligne_malformee",
                            f"{len(header)} colonnes attendues, {len(values)} trouvées",
                            raw,
                        )
                    )
                    continue
                parsed, quarantined = _convert_row(
                    company,
                    source,
                    source_type,
                    row,
                    positions,
                    raw,
                    excel_row,
                    sheet_name,
                    sheet_site,
                    sheet_unit,
                )
                if parsed is not None:
                    result.rows.append(parsed)
                if quarantined is not None:
                    result.quarantine.append(quarantined)
    except (OSError, zipfile.BadZipFile, InvalidFileException, ValueError) as exc:
        raise SourceReadError("fichier_illisible", str(exc)) from exc
    finally:
        workbook.close()
    return result


def read_source(
    company: Company,
    source: Source,
    source_type: SourceType,
    data_raw_dir: Path,
) -> ReadResult:
    """Lit une source et retourne les lignes valides et la quarantaine."""
    path = data_raw_dir / source.path
    if source.type == "electricity_xlsx":
        return _read_xlsx(path, company, source, source_type)
    return _read_delimited(path, company, source, source_type)


def _table_columns(table: str) -> list[str]:
    return [name for name, _ in RAW_COLUMNS[table]]


def _insert_expressions(columns: list[str]) -> str:
    return ", ".join(
        "CAST(? AS JSON)" if column == "raw_row" else "?" for column in columns
    )


INSERT_CHUNK_SIZE = 500


def _bulk_insert(
    con: duckdb.DuckDBPyConnection,
    table: str,
    columns: list[str],
    payload: list[tuple[Any, ...]],
) -> None:
    """Insertion multi-lignes : plus rapide que executemany sur DuckDB."""
    if not payload:
        return
    row_expression = f"({_insert_expressions(columns)})"
    for start in range(0, len(payload), INSERT_CHUNK_SIZE):
        chunk = payload[start : start + INSERT_CHUNK_SIZE]
        sql = (
            f"INSERT INTO raw.{table} ({', '.join(columns)}) VALUES "
            + ", ".join([row_expression] * len(chunk))
        )
        con.execute(sql, [value for row in chunk for value in row])


def _ensure_schema(con: duckdb.DuckDBPyConnection, rebuild: bool) -> None:
    con.execute("CREATE SCHEMA IF NOT EXISTS raw")
    if rebuild:
        for table in sorted(set(TABLE_BY_TYPE.values())):
            con.execute(f"DROP TABLE IF EXISTS raw.{table}")

    for table, columns in RAW_COLUMNS.items():
        definitions = [f"{name} {sql_type}" for name, sql_type in columns + META_COLUMNS]
        con.execute(
            f"CREATE TABLE IF NOT EXISTS raw.{table} ({', '.join(definitions)})"
        )
        expected = [name for name, _ in columns] + [name for name, _ in META_COLUMNS]
        actual = [
            row[1] for row in con.execute(f"PRAGMA table_info('raw.{table}')").fetchall()
        ]
        if actual != expected:
            raise IngestError(
                f"Schéma obsolète pour raw.{table}: "
                f"attendu {expected}, trouvé {actual}. Relancer avec --rebuild."
            )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS raw.ingestion_log (
            ingestion_id VARCHAR NOT NULL,
            company_code VARCHAR NOT NULL,
            source_id VARCHAR NOT NULL,
            source_type VARCHAR NOT NULL,
            source_file VARCHAR NOT NULL,
            source_sheet VARCHAR,
            source_sha256 VARCHAR,
            file_size_bytes BIGINT,
            rows_read INTEGER NOT NULL,
            rows_loaded INTEGER NOT NULL,
            rows_quarantined INTEGER NOT NULL,
            status VARCHAR NOT NULL,
            error VARCHAR,
            duration_ms INTEGER NOT NULL,
            started_at TIMESTAMPTZ NOT NULL,
            finished_at TIMESTAMPTZ NOT NULL
        )
        """
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS raw.quarantine (
            ingestion_id VARCHAR NOT NULL,
            company_code VARCHAR NOT NULL,
            source_id VARCHAR NOT NULL,
            source_file VARCHAR NOT NULL,
            source_sheet VARCHAR,
            source_row INTEGER,
            reason_code VARCHAR NOT NULL,
            reason_detail VARCHAR,
            raw_row JSON,
            quarantined_at TIMESTAMPTZ NOT NULL
        )
        """
    )


def _json_row(raw: dict[str, Any] | None) -> str | None:
    if raw is None:
        return None
    return json.dumps(raw, ensure_ascii=False, default=str)


def _write_rows(
    con: duckdb.DuckDBPyConnection,
    log: SourceLog,
    result: ReadResult,
    replace: bool = True,
) -> None:
    table = TABLE_BY_TYPE[log.source_type]
    data_columns = _table_columns(table)
    meta_names = [name for name, _ in META_COLUMNS]
    columns = data_columns + meta_names

    payload = []
    for row in result.rows:
        arguments = [row.values.get(name) for name in data_columns]
        arguments.extend(
            [
                log.company_code,
                log.source_id,
                log.source_file,
                row.source_sheet,
                row.source_row,
                log.ingestion_id,
                log.finished_at,
                _json_row(row.raw),
            ]
        )
        payload.append(tuple(arguments))

    quarantine_columns = [
        "ingestion_id",
        "company_code",
        "source_id",
        "source_file",
        "source_sheet",
        "source_row",
        "reason_code",
        "reason_detail",
        "raw_row",
        "quarantined_at",
    ]
    quarantine_payload = [
        (
            log.ingestion_id,
            log.company_code,
            log.source_id,
            log.source_file,
            row.source_sheet,
            row.source_row,
            row.reason_code,
            row.reason_detail,
            _json_row(row.raw),
            log.finished_at,
        )
        for row in result.quarantine
    ]

    con.execute("BEGIN")
    try:
        if replace:
            con.execute(
                f"DELETE FROM raw.{table} WHERE company_code = ? AND source_id = ?",
                [log.company_code, log.source_id],
            )
            _bulk_insert(con, table, columns, payload)
        _bulk_insert(con, "quarantine", quarantine_columns, quarantine_payload)
        con.execute("COMMIT")
    except Exception:
        with contextlib.suppress(duckdb.Error):
            con.execute("ROLLBACK")
        raise


def _write_log(con: duckdb.DuckDBPyConnection, log: SourceLog) -> None:
    con.execute(
        """
        INSERT INTO raw.ingestion_log (
            ingestion_id, company_code, source_id, source_type, source_file,
            source_sheet, source_sha256, file_size_bytes, rows_read,
            rows_loaded, rows_quarantined, status, error, duration_ms,
            started_at, finished_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            log.ingestion_id,
            log.company_code,
            log.source_id,
            log.source_type,
            log.source_file,
            log.source_sheet,
            log.source_sha256,
            log.file_size_bytes,
            log.rows_read,
            log.rows_loaded,
            log.rows_quarantined,
            log.status,
            log.error,
            log.duration_ms,
            log.started_at,
            log.finished_at,
        ],
    )


def _ingest_source(
    con: duckdb.DuckDBPyConnection,
    company: Company,
    source: Source,
    source_type: SourceType,
    data_raw_dir: Path,
) -> SourceLog:
    started_at = _now()
    started = time.perf_counter()
    ingestion_id = uuid.uuid4().hex
    path = data_raw_dir / source.path

    sha256: str | None = None
    file_size: int | None = None
    result: ReadResult
    error: str | None = None

    try:
        sha256, file_size = _hash_file(path)
    except OSError as exc:
        error = f"fichier_illisible: {exc}"
        result = ReadResult(
            quarantine=[
                QuarantineRow(None, None, "fichier_illisible", str(exc), None)
            ]
        )
    else:
        try:
            result = read_source(company, source, source_type, data_raw_dir)
        except SourceReadError as exc:
            error = f"{exc.reason_code}: {exc.detail}"
            result = ReadResult(
                quarantine=[
                    QuarantineRow(None, None, exc.reason_code, exc.detail, None)
                ]
            )

    if error is None and result.rows_read == 0:
        error = "aucune ligne de données"
    if error is None and not result.rows:
        error = "aucune ligne chargeable"

    finished_at = _now()
    duration_ms = max(1, int((time.perf_counter() - started) * 1000))
    log = SourceLog(
        ingestion_id=ingestion_id,
        company_code=company.code,
        source_id=source.id,
        source_type=source.type,
        source_file=source.path,
        source_sheet=", ".join(source.sheets) if source.sheets else None,
        source_sha256=sha256,
        file_size_bytes=file_size,
        rows_read=result.rows_read,
        rows_loaded=len(result.rows),
        rows_quarantined=len(result.quarantine),
        status="failed" if error else ("partial" if result.quarantine else "success"),
        error=error,
        duration_ms=duration_ms,
        started_at=started_at,
        finished_at=finished_at,
    )

    if error is None:
        try:
            _write_rows(con, log, result)
        except Exception as exc:
            log.status = "failed"
            log.error = f"ecriture: {exc}"
            log.rows_loaded = 0
            log.finished_at = _now()
            log.duration_ms = max(1, int((time.perf_counter() - started) * 1000))
            _write_log(con, log)
            return log
    elif result.quarantine:
        _write_rows(con, log, ReadResult(quarantine=result.quarantine), replace=False)

    log.finished_at = _now()
    log.duration_ms = max(1, int((time.perf_counter() - started) * 1000))
    _write_log(con, log)
    return log


def ingest(
    config: Config,
    data_raw_dir: Path | None = None,
    db_path: Path | None = None,
    company_codes: list[str] | None = None,
    rebuild: bool = False,
) -> RunSummary:
    """Ingère les sources sélectionnées et retourne le résumé d'exécution."""
    data_raw_dir = Path(data_raw_dir or DATA_RAW_DIR)
    db_path = Path(db_path or DEFAULT_DB_PATH)

    if company_codes:
        unknown = sorted(set(company_codes) - set(config.companies))
        if unknown:
            raise ConfigError(f"Entreprise inconnue: {', '.join(unknown)}")
        companies = [config.companies[code] for code in company_codes]
    else:
        companies = [config.companies[code] for code in sorted(config.companies)]

    db_path.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(str(db_path))
    try:
        _ensure_schema(con, rebuild)
        entries = []
        for company in companies:
            for source in company.sources:
                entries.append(
                    _ingest_source(
                        con,
                        company,
                        source,
                        config.source_types[source.type],
                        data_raw_dir,
                    )
                )
    finally:
        con.close()

    return RunSummary(
        companies=len(companies),
        files=len(entries),
        rows_loaded=sum(entry.rows_loaded for entry in entries),
        rows_quarantined=sum(entry.rows_quarantined for entry in entries),
        failures=sum(entry.status == "failed" for entry in entries),
        db_path=db_path,
        entries=entries,
    )


def format_summary(summary: RunSummary) -> str:
    lines = [
        "Ingestion terminée",
        f"Entreprises : {summary.companies}",
        f"Fichiers    : {summary.files}",
        f"Lignes      : {summary.rows_loaded}",
        f"Quarantaine : {summary.rows_quarantined}",
        f"Base        : {summary.db_path}",
    ]
    if summary.failures:
        lines.append(f"Échecs      : {summary.failures}")
        for entry in summary.entries:
            if entry.status == "failed":
                lines.append(
                    f"  - {entry.company_code}/{entry.source_id}: {entry.error}"
                )
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Ingère les fichiers sources BEGES dans DuckDB"
    )
    selection = parser.add_mutually_exclusive_group(required=True)
    selection.add_argument(
        "--all", action="store_true", help="Ingère toutes les entreprises"
    )
    selection.add_argument(
        "--company",
        action="append",
        metavar="CODE",
        help="Ingère une entreprise (répétable)",
    )
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--config-dir", type=Path, default=None)
    parser.add_argument("--data-raw-dir", type=Path, default=None)
    parser.add_argument(
        "--rebuild",
        action="store_true",
        help="Recrée les tables raw de données (journal et quarantaine conservés)",
    )
    args = parser.parse_args(argv)

    try:
        from beges.config import CONFIG_DIR

        config = load_config(
            args.config_dir or CONFIG_DIR,
            args.data_raw_dir or DATA_RAW_DIR,
        )
        summary = ingest(
            config,
            data_raw_dir=args.data_raw_dir,
            db_path=args.db,
            company_codes=None if args.all else args.company,
            rebuild=args.rebuild,
        )
    except (ConfigError, IngestError, duckdb.Error) as exc:
        print(f"ERREUR: {exc}", file=sys.stderr)
        return 1

    print(format_summary(summary))
    return 1 if summary.failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
